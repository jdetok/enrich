from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
# utf-8-sig handles a BOM (common in Excel exports); latin-1 never fails to decode.
FALLBACK_ENCODINGS = ("utf-8-sig", "latin-1")

Row = dict[str, Any]

def read_table(
    path: str | Path,
    *,
    sheet: str | None = None,
    encoding: str | None = None,
    delimiter: str | None = None,
) -> list[Row]:
    """Read a CSV/TSV or Excel file into a list of dicts (one per row).

    The first row is used as the header. The format is chosen from the file
    extension.

    CSV: the delimiter is sniffed unless given, and the encoding falls back
    from UTF-8 (with BOM) to Latin-1 unless given. Values are strings.
    Excel: reads the first sheet unless `sheet` names one. Values keep their
    native types (int, float, datetime, ...)
    """
    path = Path(path).expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"No such file: {path}")

    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path, sheet)
    if suffix in CSV_SUFFIXES:
        if delimiter is None and suffix == ".tsv":
            delimiter = "\t"
        return _read_csv(path, encoding, delimiter)

    supported = sorted(CSV_SUFFIXES | EXCEL_SUFFIXES)
    raise ValueError(f"Unsupported file type {suffix!r}. Supported: {supported}")


def _read_csv(path: Path, encoding: str | None, delimiter: str | None) -> list[Row]:
    for enc in (encoding,) if encoding else FALLBACK_ENCODINGS:
        try:
            with path.open(newline="", encoding=enc) as f:
                if delimiter is None:
                    sample = f.read(8192)
                    f.seek(0)
                    try:
                        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                    except csv.Error:
                        delimiter = ","
                return list(csv.DictReader(f, delimiter=delimiter))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode {path} as {encoding!r}")


def _read_excel(path: Path, sheet: str | None) -> list[Row]:
    from openpyxl import load_workbook  # imported lazily so CSV use needs no extras

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet is not None else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        keys = [str(h).strip() if h is not None else f"column_{i}"
                for i, h in enumerate(header, 1)]
        return [dict(zip(keys, row)) for row in rows
                if any(cell is not None for cell in row)]
    finally:
        wb.close()